
import base64
import io
import unicodedata
import re
from difflib import SequenceMatcher
import openpyxl
from odoo import api, fields, models, _
from odoo.exceptions import UserError

def normalize_name(name):
    if not name:
        return ""
    name = str(name).replace('\xa0', ' ').replace('\u200b', ' ').replace('\t', ' ')
    name = ''.join(c for c in unicodedata.normalize('NFD', name) if unicodedata.category(c) != 'Mn')
    name = name.upper()
    name = re.sub(r'[^A-Z0-9\s]', ' ', name)
    name = re.sub(r'\s+', ' ', name).strip()
    return name

def token_sort_ratio(a, b):
    a_tokens = sorted(normalize_name(a).split())
    b_tokens = sorted(normalize_name(b).split())
    return SequenceMatcher(None, ' '.join(a_tokens), ' '.join(b_tokens)).ratio()

class SgsPerdiemDepositImportWizard(models.TransientModel):
    _name = 'sgs.perdiem.deposit.import.wizard'
    _description = 'Importador inteligente de depositos'

    file = fields.Binary('Archivo Excel', required=True)
    filename = fields.Char('Nombre archivo')
    date_default = fields.Date('Fecha por defecto', default=fields.Date.context_today)
    line_ids = fields.One2many('sgs.perdiem.deposit.import.line', 'wizard_id', string='Lineas a conciliar')
    state = fields.Selection([('draft','Carga'),('to_conciliate','Conciliar'),('done','Hecho')], default='draft')

    def action_parse_file(self):
        self.ensure_one()
        data = base64.b64decode(self.file)
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        ws = wb.active
        col_map = {}
        header_row_idx = 0
        is_bank_format = False

        # Detecta formato por header
        for r in range(1, 20):
            row_vals = [str(c.value or '') for c in ws[r]]
            norm_row = [normalize_name(v) for v in row_vals]
            if not any(row_vals):
                continue
            # Formato viaticos clasico
            if any('CUSTODIO' in n for n in norm_row):
                header_row_idx = r
                for idx, h in enumerate(norm_row):
                    if 'FECHA' in h and 'DEPOSITO' in h: col_map['date']=idx
                    elif 'CUSTODIO' in h: col_map['custodio']=idx
                    elif 'MONTO' in h or 'IMPORTE' in h: col_map['amount']=idx
                    elif 'CONCEPTO' in h: col_map['concept']=idx
                break
            # Formato PLANTILLA_SISTEMAS_GLOBALES (tu archivo actual)
            if 'IUT' in norm_row and 'MONTO TRASPASO' in ' '.join(norm_row) and 'EMPLEADO' in ' '.join(norm_row):
                header_row_idx = r
                # En tu plantilla: IUT=0, MONTO=1, EMPLEADO=2, N EMPLEADO=3
                for idx, h in enumerate(norm_row):
                    if h == 'IUT': col_map['iut']=idx
                    elif 'MONTO' in h: col_map['amount']=idx
                    elif h == 'EMPLEADO': col_map['custodio']=idx
                    elif 'N DE EMPLEADO' in h or 'NUM' in h: col_map['num_emp']=idx
                is_bank_format = True
                break

        if not col_map.get('custodio') and not col_map.get('amount'):
            # fallback: si no encontro header pero tiene datos tipo banco
            col_map = {'iut':0, 'amount':1, 'custodio':2, 'num_emp':3}
            header_row_idx = 6  # tu plantilla tiene header en fila 6
            is_bank_format = True

        custodians = self.env['sgs.custodian'].search([])
        employees = self.env['hr.employee'].search([])
        cust_map = {normalize_name(c.name): c for c in custodians}
        emp_map = {normalize_name(e.name): e for e in employees}

        lines = []
        skipped_zero = 0
        for row in ws.iter_rows(min_row=header_row_idx+1, values_only=True):
            if not row or not any(row):
                continue
            row = list(row) + [None]*10
            raw_cust = str(row[col_map.get('custodio',2)] or '').strip()
            if not raw_cust: continue
            if normalize_name(raw_cust) in ('EMPLEADO','CUSTODIO') or len(raw_cust) < 4:
                continue
            if raw_cust.replace('.','',1).isdigit():
                continue

            raw_amount = row[col_map.get('amount',1)]
            try:
                amount = float(str(raw_amount).replace(',','').replace('$','').replace(' ','') or 0)
            except:
                amount = 0

            if amount == 0:
                skipped_zero += 1
                continue  # tu plantilla trae muchos en 0, los ignoramos como pediste

            raw_date = row[col_map.get('date',0)] if col_map.get('date') is not None else None
            date_val = self.date_default
            if raw_date and hasattr(raw_date, 'year'):
                date_val = raw_date.date() if hasattr(raw_date, 'date') else raw_date

            iut = str(row[col_map.get('iut',0)] or '').strip() if 'iut' in col_map else ''
            num_emp = str(row[col_map.get('num_emp',3)] or '').strip() if 'num_emp' in col_map else ''
            concept = f"{iut} - Emp {num_emp}".strip() if is_bank_format else str(row[col_map.get('concept',4)] or 'Deposito viaticos').strip()

            norm = normalize_name(raw_cust)
            custodian = cust_map.get(norm)
            employee = emp_map.get(norm)
            if not custodian and employee:
                custodian = self.env['sgs.custodian'].search([('employee_id','=',employee.id)], limit=1)
            
            status = 'matched'
            score = 1.0
            if not custodian:
                best_match = None
                best_score = 0
                for c_name, c_rec in cust_map.items():
                    if not set(norm.split()) & set(c_name.split()): continue
                    s = token_sort_ratio(norm, c_name)
                    if s > best_score:
                        best_score = s
                        best_match = c_rec
                if best_score >= 0.80:
                    custodian = best_match; score=best_score; status='matched_auto'
                elif best_score >= 0.45:
                    custodian = best_match; score=best_score; status='to_conciliate'
                else:
                    status='not_found'; score=best_score

            lines.append((0,0,{
                'custodio_raw': raw_cust,
                'custodio_normalized': norm,
                'custodian_id': custodian.id if custodian else False,
                'employee_id': employee.id if employee else (custodian.employee_id.id if custodian and custodian.employee_id else False),
                'date': date_val,
                'amount': amount,
                'concept': concept[:200],
                'match_score': score,
                'state': status,
                'to_import': bool(custodian and amount > 0),
            }))

        if not lines:
            raise UserError(f"No se detectaron lineas validas con monto >0. Se omitieron {skipped_zero} lineas con monto 0. Mapa: {col_map} header fila {header_row_idx}. Tu plantilla debe tener MONTO TRASPASO >0 en la columna B.")

        self.line_ids = [(5,0,0)] + lines
        self.state = 'to_conciliate'
        return {
            'type':'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode':'form',
            'target':'new',
        }

    def action_create_deposits(self):
        to_create = self.line_ids.filtered(lambda l: l.custodian_id and l.to_import and l.amount > 0)
        if not to_create:
            raise UserError("No hay lineas validas para importar.")
        deposits = self.env['sgs.perdiem.deposit'].create([{
            'custodian_id': l.custodian_id.id,
            'date': l.date,
            'amount': l.amount,
            'concept': l.concept,
        } for l in to_create])
        self.state = 'done'
        return {
            'type':'ir.actions.act_window',
            'name':'Depositos creados',
            'res_model':'sgs.perdiem.deposit',
            'domain':[('id','in',deposits.ids)],
            'view_mode':'list,form',
        }

class SgsPerdiemDepositImportLine(models.TransientModel):
    _name = 'sgs.perdiem.deposit.import.line'
    _description = 'Linea de conciliacion de deposito'
    wizard_id = fields.Many2one('sgs.perdiem.deposit.import.wizard', required=True, ondelete='cascade')
    custodio_raw = fields.Char('Nombre en Excel', readonly=True)
    custodio_normalized = fields.Char('Normalizado', readonly=True)
    employee_id = fields.Many2one('hr.employee', string='Empleado encontrado')
    custodian_id = fields.Many2one('sgs.custodian', string='Custodio a depositar')
    date = fields.Date('Fecha')
    amount = fields.Float('Monto')
    concept = fields.Char('Concepto')
    match_score = fields.Float('Score', digits=(3,2))
    state = fields.Selection([
        ('matched','Coincidencia exacta'),
        ('matched_auto','Auto-conciliado'),
        ('to_conciliate','Requiere revision'),
        ('not_found','No encontrado')
    ], default='to_conciliate')
    to_import = fields.Boolean('Importar', default=True)
